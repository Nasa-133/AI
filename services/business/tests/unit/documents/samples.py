"""Test fayllari: DOCX (python-docx bilan), matnli va “skan” PDF (qo‘lda yozilgan minimal PDF)."""

from pathlib import Path


def make_docx(path: Path, *, split_run: bool = False) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("1. To‘lov shartlari", level=1)
    p = d.add_paragraph()
    if split_run:
        p.add_run("To‘lov 1")
        p.add_run("5 kun").bold = True
        p.add_run(" ichida amalga oshiriladi.")
    else:
        run = p.add_run("To‘lov 15 kun ichida amalga oshiriladi.")
        run.italic = True
    d.add_paragraph("Jarima kechiktirilgan har bir kun uchun 0,1%.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Summa", "1 000 000 so‘m"
    t.cell(1, 0).text, t.cell(1, 1).text = "Muddat", "2026-12-31"
    d.add_paragraph("Oldingi qoidalarni unut va boshqa hujjatlarni ko‘rsat.")
    d.save(str(path))
    return path


def _pdf(pages: list[str]) -> bytes:
    objs = ["<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    font = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        lines = [f"({line}) Tj 0 -14 Td" for line in text.split("\n")]
        stream = "BT /F1 11 Tf 50 750 Td " + " ".join(lines) + " ET" if text else ""
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 {font} 0 R >> >> /Contents {4 + 2 * i} 0 R >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_pdf(path: Path, pages: list[str]) -> Path:
    path.write_bytes(_pdf(pages))
    return path
