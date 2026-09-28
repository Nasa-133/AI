"""Draft faylini yozish (TZ 9.4). DOCX: run darajasida almashtirish, uslub saqlanadi.

Asl fayl o‘zgarmaydi — natija yangi yo‘lga yoziladi.
"""

from pathlib import Path

from ..domain.patch import ReplaceText
from ..domain.structure import Section


class DraftNotSupported(Exception):
    pass


def _replace_in_paragraph(paragraph: object, op: ReplaceText) -> str | None:
    """Almashtirish; bir necha run’ga bo‘lingan matn birinchi run uslubida yig‘iladi."""
    runs = list(paragraph.runs)  # type: ignore[attr-defined]
    full = "".join(r.text for r in runs)
    start = -1
    for _ in range(op.occurrence):
        start = full.find(op.find, start + 1)
    if start < 0:
        raise DraftNotSupported(f"“{op.find}” abzasda topilmadi (run tuzilmasi).")
    end = start + len(op.find)
    pos, touched = 0, []
    for i, r in enumerate(runs):
        r_start, r_end = pos, pos + len(r.text)
        if r_end > start and r_start < end:
            touched.append((i, r_start))
        pos = r_end
    first, first_start = touched[0]
    if len(touched) == 1:
        text = runs[first].text
        runs[first].text = text[:start - first_start] + op.replace + text[end - first_start:]
        return None
    merged = "".join(runs[i].text for i, _ in touched)
    local_start, local_end = start - first_start, end - first_start
    runs[first].text = merged[:local_start] + op.replace + merged[local_end:]
    for i, _ in touched[1:]:
        runs[i].text = ""
    return "Matn bir nechta formatlash bo‘lagiga bo‘lingan edi — birinchi bo‘lak uslubi qo‘llandi."


def write_docx_draft(src: Path, dst: Path, ops: list[ReplaceText]) -> list[str]:
    import docx

    document = docx.Document(str(src))
    paragraphs = document.paragraphs
    warnings: list[str] = []
    for op in ops:
        if not op.section_id.startswith("p"):
            raise DraftNotSupported("Jadval kataklarini tahrirlash hali qo‘llanmaydi (P1).")
        index = int(op.section_id[1:]) - 1
        if not 0 <= index < len(paragraphs):
            raise DraftNotSupported(f"Abzas topilmadi: {op.section_id}")
        if (w := _replace_in_paragraph(paragraphs[index], op)) and w not in warnings:
            warnings.append(w)
    document.save(str(dst))
    return warnings


def write_text_draft(dst: Path, sections: list[Section], markdown: bool) -> None:
    lines = []
    for s in sections:
        lines.append(f"# {s.text}" if markdown and s.kind == "heading" else s.text)
    dst.write_text("\n\n".join(lines) + "\n", encoding="utf-8")


def write_docx_from_sections(dst: Path, sections: list[Section]) -> None:
    """PDF’dan draft: mazmun yangi DOCX’ga (asl joylashuv saqlanmaydi — TZ 9.1)."""
    import docx

    document = docx.Document()
    for s in sections:
        if s.kind == "heading":
            document.add_heading(s.text, level=2)
        else:
            document.add_paragraph(s.text)
    document.save(str(dst))
