"""PDF/DOCX/TXT → bo‘limlar. Alohida jarayonda chaqiriladi (sandbox.py); tarmoq va DB yo‘q.

Bo‘lim ID’lari manbaga bog‘langan va barqaror: DOCX `p{abzas}` / `t{jadval}r{qator}`,
PDF `pg{sahifa}p{abzas}`, TXT/MD `s{tartib}`. DOCX patch aynan shu abzas raqamini ishlatadi.
"""

import re
from pathlib import Path
from typing import Any

MIN_PAGE_TEXT = 20
PDF_NEEDS_OCR_RATIO = 0.5


class InvalidDocument(Exception):
    """Faylni o‘qib bo‘lmaydi (parol, buzilgan, qo‘llanmaydigan tuzilma) — aniq sabab bilan."""


def _section(sid: str, ordinal: int, kind: str, locator: str, text: str) -> dict[str, Any]:
    return {"section_id": sid, "ordinal": ordinal, "kind": kind, "locator": locator, "text": text}


def parse_pdf(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise InvalidDocument("PDF parol bilan himoyalangan. Parolsiz nusxasini yuklang.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise InvalidDocument(f"PDF buzilgan yoki o‘qilmaydi: {exc}") from exc
    sections: list[dict[str, Any]] = []
    empty = 0
    for p, text in enumerate(pages, start=1):
        paragraphs = [re.sub(r"\s*\n\s*", " ", x).strip() for x in re.split(r"\n\s*\n", text)]
        paragraphs = [x for x in paragraphs if x]
        if sum(len(x) for x in paragraphs) < MIN_PAGE_TEXT:
            empty += 1
            continue
        for k, para in enumerate(paragraphs, start=1):
            sections.append(_section(f"pg{p}p{k}", len(sections), "paragraph", f"sahifa {p}, ¶{k}",
                                     para))
    needs_ocr = not pages or empty / len(pages) >= PDF_NEEDS_OCR_RATIO
    return {"sections": sections, "quality": {
        "pages": len(pages), "empty_pages": empty, "needs_ocr": needs_ocr,
        "warnings": (["Sahifalarning ko‘pida matn yo‘q — skan PDF bo‘lishi mumkin; OCR P1 da."]
                     if needs_ocr else [])}}


def parse_docx(path: Path) -> dict[str, Any]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise InvalidDocument(f"DOCX buzilgan yoki o‘qilmaydi: {exc}") from exc
    sections: list[dict[str, Any]] = []
    para_no = table_no = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            para_no += 1
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "") or ""
            heading = style.lower().startswith(("heading", "title", "sarlavha"))
            kind = "heading" if heading else "paragraph"
            sections.append(_section(f"p{para_no}", len(sections), kind, f"¶{para_no}", text))
        elif isinstance(block, Table):
            table_no += 1
            for r, row in enumerate(block.rows, start=1):
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    sections.append(_section(f"t{table_no}r{r}", len(sections), "table_row",
                                             f"jadval {table_no}, qator {r}", " | ".join(cells)))
    xml = document.element.xml
    warnings = []
    if "w:txbxContent" in xml:
        warnings.append("Hujjatda matn qutilari bor — ular o‘qilmadi va tahrirlanmaydi.")
    if "w:fldSimple" in xml or "w:instrText" in xml:
        warnings.append("Hujjatda avtomatik maydonlar bor — tahrirda o‘zgartirilmaydi.")
    if "w:object" in xml or "o:OLEObject" in xml:
        warnings.append("Hujjatda joylashtirilgan obyektlar bor — ular o‘qilmadi.")
    return {"sections": sections, "quality": {"paragraphs": para_no, "tables": table_no,
                                              "needs_ocr": False, "warnings": warnings}}


def parse_text(path: Path) -> dict[str, Any]:
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise InvalidDocument("Matn UTF-8 kodlashda emas. Faylni UTF-8 sifatida saqlang.") from exc
    sections: list[dict[str, Any]] = []
    heading, para = "", 0
    for block in re.split(r"\n\s*\n", raw):
        text = block.strip()
        if not text:
            continue
        if m := re.match(r"^#{1,6}\s+(.*)", text):
            heading, para = m.group(1).strip(), 0
            sections.append(_section(f"s{len(sections)}", len(sections), "heading",
                                     f"§ {heading}", heading))
            continue
        para += 1
        loc = f"§ {heading} › ¶{para}" if heading else f"¶{para}"
        sections.append(_section(f"s{len(sections)}", len(sections), "paragraph", loc,
                                 re.sub(r"\s*\n\s*", " ", text)))
    return {"sections": sections, "quality": {"needs_ocr": False, "warnings": []}}


PARSERS = {"pdf": parse_pdf, "docx": parse_docx, "txt": parse_text, "md": parse_text}
